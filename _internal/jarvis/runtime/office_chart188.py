"""NEO 188 NEO-OFFICE-001 / TZ-OFFICE-EDITABLE-CHART-01.

«В файле Продажи.xlsx, лист Итоги, по A1:B13 создай столбчатую диаграмму в
новом файле» → a preview (exact source, sheet, range, headers, sample values,
chart type, new file name, «копия выбранных данных; исходник не изменяется»)
→ after «да» a NEW ordinary .xlsx with a copy of only those values and a real,
editable bar/line chart object. The source is never written.

First version, deliberately narrow: a saved .xlsx (not .xlsm/.xls/.csv), one
visible sheet, one bounded A1 range with a header row; first column =
categories, next 1–5 columns = numeric series. Formulas, booleans, errors,
text numbers, blanks, merged or hidden cells are refused with the exact cell —
nothing is guessed, summed or recalculated. No Office, macros, external links
or generated code are run. The new file never overwrites an existing one.
"""
from __future__ import annotations

import hashlib
import math
import os
import re
import uuid
import zipfile
from functools import lru_cache
from pathlib import Path
from typing import Any

from jarvis.runtime.natural_commands import LANGUAGES

MAX_ROWS = 2000
MAX_SERIES = 5
_RANGE = re.compile(r'\b(?P<a>[A-Za-z]{1,3})(?P<r1>\d{1,7})\s*:\s*(?P<b>[A-Za-z]{1,3})(?P<r2>\d{1,7})\b')
_FILE = re.compile(r'(?:«(?P<q>[^»]+?\.xlsx)»|"(?P<d>[^"]+?\.xlsx)"|(?P<p>[^\s,;:«»"]+?\.xlsx))', re.I)
_SHEET = {
    'ru': r'(?:на\s+)?лист[еа]?\s+(?P<s>«[^»]+»|"[^"]+"|[^\s,;:]+)', 'en': r'(?:on\s+)?(?:the\s+)?sheet\s+(?P<s>«[^»]+»|"[^"]+"|[^\s,;:]+)',
    'de': r'(?:im\s+)?blatt\s+(?P<s>«[^»]+»|"[^"]+"|[^\s,;:]+)', 'fr': r'(?:de\s+la\s+|la\s+)?feuille\s+(?P<s>«[^»]+»|"[^"]+"|[^\s,;:]+)',
    'es': r'(?:de\s+la\s+|la\s+)?hoja\s+(?P<s>«[^»]+»|"[^"]+"|[^\s,;:]+)', 'it': r'(?:del\s+)?foglio\s+(?P<s>«[^»]+»|"[^"]+"|[^\s,;:]+)',
    'pt': r'(?:da\s+)?folha\s+(?P<s>«[^»]+»|"[^"]+"|[^\s,;:]+)',
}
_CHART = {
    'ru': (r'диаграмм|график', r'столбчат|гистограмм|столбц', r'линейн|график'),
    'en': (r'chart|graph', r'bar|column', r'line'),
    'de': (r'diagramm', r'säulen|saeulen|balken', r'linien'),
    'fr': (r'graphique|diagramme', r'colonnes|barres|histogramme', r'ligne|courbe'),
    'es': (r'gr[aá]fic[oa]|diagrama', r'columnas|barras', r'l[ií]neas?'),
    'it': (r'grafico|diagramma', r'colonne|barre|istogramma', r'linee|linea'),
    'pt': (r'gr[aá]fico|diagrama', r'colunas|barras', r'linhas?'),
}
_CREATE = {'ru': r'создай|построй|сделай|нарисуй', 'en': r'create|make|build|draw', 'de': r'erstelle|mache|zeichne',
           'fr': r'crée|cree|fais|dessine', 'es': r'crea|haz|dibuja', 'it': r'crea|fai|disegna', 'pt': r'cria|crie|faz|desenha'}


def _unquote(value: str) -> str:
    return str(value or '').strip().strip('«»"“”').strip()


def parse(text: str, code: str) -> tuple[str, dict] | None:
    """A complete or partial chart request (the tool asks for anything missing)."""
    if code not in LANGUAGES:
        return None
    value = str(text or '')
    low = value.casefold()
    noun, bar, line = _CHART[code]
    if not re.search(r'(?:' + _CREATE[code] + r')', low) or not re.search(noun, low):
        return None
    file = _FILE.search(value)
    rng = _RANGE.search(value)
    if file is None and rng is None:
        return None                   # an ordinary «create a chart» stays with the existing Office route
    args: dict[str, Any] = {'action': 'chart'}
    if file:
        args['name'] = _unquote(file.group('q') or file.group('d') or file.group('p'))
    if rng:
        args['range'] = (rng.group('a') + rng.group('r1') + ':' + rng.group('b') + rng.group('r2')).upper()
    sheet = re.search(_SHEET[code], value, re.I)
    if sheet:
        args['sheet'] = _unquote(sheet.group('s'))
    if re.search(bar, low):
        args['chart'] = 'bar'
    elif re.search(line, low) and not (code == 'ru' and 'график' in low and not re.search(r'линейн', low)):
        args['chart'] = 'line'
    return 'office', args


_ORDER = ('ru', 'en', 'de', 'fr', 'es', 'it', 'pt')
_TEXT = {
    'need': ('Для диаграммы не хватает: {what}. Например: «в файле Продажи.xlsx, лист Итоги, по A1:B13 создай столбчатую диаграмму».',
             'The chart still needs: {what}. For example: “create a bar chart from A1:B13 on sheet Summary in Sales.xlsx”.',
             'Für das Diagramm fehlt: {what}. Beispiel: „Erstelle aus A1:B13 im Blatt Übersicht der Datei Umsatz.xlsx ein Säulendiagramm“.',
             'Il manque pour le graphique : {what}. Par ex. « crée un graphique en colonnes à partir de A1:B13, feuille Résumé, fichier Ventes.xlsx ».',
             'Para el gráfico falta: {what}. Por ejemplo: «crea un gráfico de columnas con A1:B13 de la hoja Resumen de Ventas.xlsx».',
             'Per il grafico manca: {what}. Per esempio: «crea un grafico a colonne usando A1:B13 del foglio Riepilogo di Vendite.xlsx».',
             'Para o gráfico falta: {what}. Por exemplo: «cria um gráfico de colunas com A1:B13 da folha Resumo de Vendas.xlsx».'),
    'what': (('файл .xlsx', 'лист', 'диапазон вроде A1:B13', 'тип: столбчатая или линейная'),
             ('the .xlsx file', 'the sheet', 'a range such as A1:B13', 'the type: bar or line'),
             ('die .xlsx-Datei', 'das Blatt', 'ein Bereich wie A1:B13', 'der Typ: Säulen oder Linien'),
             ('le fichier .xlsx', 'la feuille', 'une plage comme A1:B13', 'le type : colonnes ou lignes'),
             ('el archivo .xlsx', 'la hoja', 'un rango como A1:B13', 'el tipo: columnas o líneas'),
             ('il file .xlsx', 'il foglio', 'un intervallo come A1:B13', 'il tipo: colonne o linee'),
             ('o ficheiro .xlsx', 'a folha', 'um intervalo como A1:B13', 'o tipo: colunas ou linhas')),
    'bad': ('Диаграмму не создаю: {reason}.', 'I am not creating the chart: {reason}.', 'Ich erstelle das Diagramm nicht: {reason}.',
            'Je ne crée pas le graphique : {reason}.', 'No creo el gráfico: {reason}.', 'Non creo il grafico: {reason}.', 'Não crio o gráfico: {reason}.'),
    'preview': ('Предпросмотр диаграммы:\n• Источник: {source} (лист «{sheet}», {range})\n• Подписи: столбец «{category}»; ряды: {series}\n'
                '• Строк данных: {rows}; первые значения: {sample}\n• Тип: {kind}\n• Новый файл: {destination}\n'
                'Это копия выбранных данных; исходный файл не изменяется. Создать?',
                'Chart preview:\n• Source: {source} (sheet “{sheet}”, {range})\n• Labels: column “{category}”; series: {series}\n'
                '• Data rows: {rows}; first values: {sample}\n• Type: {kind}\n• New file: {destination}\n'
                'This is a copy of the selected data; the source file is not changed. Create it?',
                'Diagrammvorschau:\n• Quelle: {source} (Blatt „{sheet}“, {range})\n• Beschriftungen: Spalte „{category}“; Reihen: {series}\n'
                '• Datenzeilen: {rows}; erste Werte: {sample}\n• Typ: {kind}\n• Neue Datei: {destination}\n'
                'Das ist eine Kopie der gewählten Daten; die Quelldatei bleibt unverändert. Erstellen?',
                'Aperçu du graphique :\n• Source : {source} (feuille « {sheet} », {range})\n• Étiquettes : colonne « {category} » ; séries : {series}\n'
                '• Lignes de données : {rows} ; premières valeurs : {sample}\n• Type : {kind}\n• Nouveau fichier : {destination}\n'
                'C’est une copie des données choisies ; le fichier source n’est pas modifié. Créer ?',
                'Vista previa del gráfico:\n• Origen: {source} (hoja «{sheet}», {range})\n• Etiquetas: columna «{category}»; series: {series}\n'
                '• Filas de datos: {rows}; primeros valores: {sample}\n• Tipo: {kind}\n• Archivo nuevo: {destination}\n'
                'Es una copia de los datos elegidos; el archivo de origen no cambia. ¿Lo creo?',
                'Anteprima del grafico:\n• Origine: {source} (foglio «{sheet}», {range})\n• Etichette: colonna «{category}»; serie: {series}\n'
                '• Righe di dati: {rows}; primi valori: {sample}\n• Tipo: {kind}\n• Nuovo file: {destination}\n'
                'È una copia dei dati scelti; il file di origine non viene modificato. Creo?',
                'Pré-visualização do gráfico:\n• Origem: {source} (folha «{sheet}», {range})\n• Rótulos: coluna «{category}»; séries: {series}\n'
                '• Linhas de dados: {rows}; primeiros valores: {sample}\n• Tipo: {kind}\n• Novo ficheiro: {destination}\n'
                'É uma cópia dos dados escolhidos; o ficheiro de origem não é alterado. Crio?'),
    'kind': (('столбчатая', 'линейная'), ('bar', 'line'), ('Säulen', 'Linien'), ('colonnes', 'lignes'), ('columnas', 'líneas'),
             ('colonne', 'linee'), ('colunas', 'linhas')),
    'done': ('Создал новый файл с редактируемой диаграммой: {destination}. Исходный файл не менял.',
             'Created a new file with an editable chart: {destination}. The source file was not changed.',
             'Neue Datei mit bearbeitbarem Diagramm erstellt: {destination}. Die Quelldatei blieb unverändert.',
             'Nouveau fichier avec un graphique modifiable créé : {destination}. Le fichier source n’a pas été modifié.',
             'Creé un archivo nuevo con un gráfico editable: {destination}. No cambié el archivo de origen.',
             'Ho creato un nuovo file con un grafico modificabile: {destination}. Il file di origine non è stato modificato.',
             'Criei um novo ficheiro com um gráfico editável: {destination}. Não alterei o ficheiro de origem.'),
    'changed': ('Исходный файл изменился после предпросмотра — ничего не создаю. Повторите команду, чтобы увидеть новый предпросмотр.',
                'The source file changed after the preview — nothing created. Repeat the request to see a new preview.',
                'Die Quelldatei hat sich nach der Vorschau geändert – nichts erstellt. Wiederhole die Anfrage für eine neue Vorschau.',
                'Le fichier source a changé après l’aperçu : rien n’est créé. Refaites la demande pour un nouvel aperçu.',
                'El archivo de origen cambió tras la vista previa: no creo nada. Repite la petición para ver una nueva.',
                'Il file di origine è cambiato dopo l’anteprima: non creo nulla. Ripeti la richiesta per una nuova anteprima.',
                'O ficheiro de origem mudou depois da pré-visualização: não crio nada. Repete o pedido para uma nova.'),
}
_REASONS = {
    'missing_file': ('файл «{name}» не найден в Документах, на Рабочем столе или в Загрузках', 'the file “{name}” was not found in Documents, Desktop or Downloads',
                     'die Datei „{name}“ wurde in Dokumente, Desktop oder Downloads nicht gefunden', 'le fichier « {name} » est introuvable dans Documents, Bureau ou Téléchargements',
                     'no encontré el archivo «{name}» en Documentos, Escritorio o Descargas', 'il file «{name}» non è in Documenti, Desktop o Download',
                     'o ficheiro «{name}» não foi encontrado em Documentos, Ambiente de trabalho ou Transferências'),
    'ambiguous': ('файлов «{name}» несколько: {paths} — укажите полный путь', 'there are several “{name}” files: {paths} — give the full path',
                  'es gibt mehrere Dateien „{name}“: {paths} – nenne den vollständigen Pfad', 'plusieurs fichiers « {name} » : {paths} — indiquez le chemin complet',
                  'hay varios archivos «{name}»: {paths}; indica la ruta completa', 'ci sono più file «{name}»: {paths}; indica il percorso completo',
                  'há vários ficheiros «{name}»: {paths}; indica o caminho completo'),
    'not_xlsx': ('поддерживается только обычный сохранённый .xlsx', 'only an ordinary saved .xlsx is supported', 'nur eine normale gespeicherte .xlsx wird unterstützt',
                 'seul un fichier .xlsx ordinaire enregistré est pris en charge', 'solo se admite un .xlsx normal guardado', 'è supportato solo un normale .xlsx salvato',
                 'só é suportado um .xlsx normal guardado'),
    'unreadable': ('файл не удалось прочитать (повреждён, зашифрован или открыт с блокировкой)', 'the file could not be read (damaged, encrypted or locked)',
                   'die Datei konnte nicht gelesen werden (beschädigt, verschlüsselt oder gesperrt)', 'le fichier n’a pas pu être lu (endommagé, chiffré ou verrouillé)',
                   'no se pudo leer el archivo (dañado, cifrado o bloqueado)', 'non è stato possibile leggere il file (danneggiato, cifrato o bloccato)',
                   'não foi possível ler o ficheiro (danificado, cifrado ou bloqueado)'),
    'no_sheet': ('листа «{sheet}» нет; есть: {sheets}', 'there is no sheet “{sheet}”; available: {sheets}', 'es gibt kein Blatt „{sheet}“; vorhanden: {sheets}',
                 'il n’y a pas de feuille « {sheet} » ; disponibles : {sheets}', 'no hay hoja «{sheet}»; disponibles: {sheets}', 'non c’è il foglio «{sheet}»; disponibili: {sheets}',
                 'não há folha «{sheet}»; disponíveis: {sheets}'),
    'hidden_sheet': ('лист «{sheet}» скрыт', 'the sheet “{sheet}” is hidden', 'das Blatt „{sheet}“ ist ausgeblendet', 'la feuille « {sheet} » est masquée',
                     'la hoja «{sheet}» está oculta', 'il foglio «{sheet}» è nascosto', 'a folha «{sheet}» está oculta'),
    'bad_range': ('диапазон {range} должен содержать строку заголовков, 1 столбец подписей, от 1 до 5 числовых столбцов и не больше 2000 строк',
                  'the range {range} needs a header row, one label column, 1 to 5 numeric columns and at most 2000 rows',
                  'der Bereich {range} braucht eine Kopfzeile, eine Beschriftungsspalte, 1–5 Zahlenspalten und höchstens 2000 Zeilen',
                  'la plage {range} doit avoir une ligne d’en-tête, une colonne d’étiquettes, 1 à 5 colonnes numériques et 2000 lignes au plus',
                  'el rango {range} necesita una fila de encabezados, una columna de etiquetas, de 1 a 5 columnas numéricas y como máximo 2000 filas',
                  'l’intervallo {range} deve avere una riga di intestazione, una colonna di etichette, da 1 a 5 colonne numeriche e al massimo 2000 righe',
                  'o intervalo {range} precisa de uma linha de cabeçalho, uma coluna de rótulos, 1 a 5 colunas numéricas e no máximo 2000 linhas'),
    'merged': ('в диапазоне есть объединённые ячейки ({cells})', 'the range contains merged cells ({cells})', 'der Bereich enthält verbundene Zellen ({cells})',
               'la plage contient des cellules fusionnées ({cells})', 'el rango tiene celdas combinadas ({cells})', 'l’intervallo contiene celle unite ({cells})',
               'o intervalo tem células unidas ({cells})'),
    'hidden': ('в диапазоне есть скрытые строки или столбцы', 'the range contains hidden rows or columns', 'der Bereich enthält ausgeblendete Zeilen oder Spalten',
               'la plage contient des lignes ou colonnes masquées', 'el rango tiene filas o columnas ocultas', 'l’intervallo contiene righe o colonne nascoste',
               'o intervalo tem linhas ou colunas ocultas'),
    'header': ('пустой заголовок в ячейке {cell}', 'empty header in cell {cell}', 'leere Überschrift in Zelle {cell}', 'en-tête vide dans la cellule {cell}',
               'encabezado vacío en la celda {cell}', 'intestazione vuota nella cella {cell}', 'cabeçalho vazio na célula {cell}'),
    'formula': ('в ячейке {cell} формула — нужен диапазон только со значениями', 'cell {cell} holds a formula — a values-only range is needed',
                'Zelle {cell} enthält eine Formel – nötig ist ein Bereich nur mit Werten', 'la cellule {cell} contient une formule : il faut une plage de valeurs seules',
                'la celda {cell} tiene una fórmula: hace falta un rango solo con valores', 'la cella {cell} contiene una formula: serve un intervallo di soli valori',
                'a célula {cell} tem uma fórmula: é preciso um intervalo só com valores'),
    'not_number': ('в ячейке {cell} не число ({value})', 'cell {cell} is not a number ({value})', 'Zelle {cell} ist keine Zahl ({value})',
                   'la cellule {cell} n’est pas un nombre ({value})', 'la celda {cell} no es un número ({value})', 'la cella {cell} non è un numero ({value})',
                   'a célula {cell} não é um número ({value})'),
    'empty': ('ячейка {cell} пуста — пропуски не заменяю нулём', 'cell {cell} is empty — gaps are not replaced with zero',
              'Zelle {cell} ist leer – Lücken werden nicht durch null ersetzt', 'la cellule {cell} est vide : les trous ne sont pas remplacés par zéro',
              'la celda {cell} está vacía: no relleno huecos con cero', 'la cella {cell} è vuota: i vuoti non vengono sostituiti con zero',
              'a célula {cell} está vazia: não substituo falhas por zero'),
    'label': ('пустая подпись в ячейке {cell}', 'empty label in cell {cell}', 'leere Beschriftung in Zelle {cell}', 'étiquette vide dans la cellule {cell}',
              'etiqueta vacía en la celda {cell}', 'etichetta vuota nella cella {cell}', 'rótulo vazio na célula {cell}'),
}


def t(key: str, code: str, **values) -> str:
    return _TEXT[key][_ORDER.index(code) if code in _ORDER else 1].format(**values)


def reason(key: str, code: str, **values) -> str:
    return _REASONS[key][_ORDER.index(code) if code in _ORDER else 1].format(**values)


def _candidates(name: str) -> list[str]:
    from jarvis.runtime.office178 import user_folders
    raw = _unquote(name)
    direct = Path(os.path.expandvars(os.path.expanduser(raw)))
    if direct.is_absolute():
        return [str(direct)] if direct.is_file() else []
    if raw.startswith(('\\\\', '//')) or '://' in raw:
        return []                                          # no URL/UNC/network locations
    found = []
    for folder in user_folders():
        try:
            items = list(folder.glob('*')) + list(folder.glob('*/*'))
        except OSError:
            continue
        found.extend(str(p) for p in items[:6000] if p.name.casefold() == raw.casefold() and p.is_file())
    return sorted(set(found))


def fingerprint(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for block in iter(lambda: handle.read(1 << 20), b''):
            digest.update(block)
    return digest.hexdigest()[:20]


def _coords(rng: str) -> tuple[int, int, int, int]:
    from openpyxl.utils.cell import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(rng)
    return min_col, min_row, max_col, max_row


def preview(config: Any, args: dict[str, Any], code: str) -> dict[str, Any]:
    """Validate the request; {'ok': True, spec...} or {'ok': False, 'text': ...}. Writes nothing."""
    missing = [label for key, label in zip(('name', 'sheet', 'range', 'chart'), _TEXT['what'][_ORDER.index(code) if code in _ORDER else 1])
               if not str(args.get(key, '') or '').strip()]
    if missing:
        return {'ok': False, 'text': t('need', code, what=', '.join(missing)), 'clarify': True}
    name = str(args['name'])
    candidates = [args['path']] if args.get('path') else _candidates(name)
    if not candidates:
        return {'ok': False, 'text': t('bad', code, reason=reason('missing_file', code, name=name))}
    if len(candidates) > 1:
        return {'ok': False, 'text': t('bad', code, reason=reason('ambiguous', code, name=name, paths='; '.join(candidates[:5]))), 'clarify': True}
    path = Path(candidates[0])
    if path.suffix.casefold() != '.xlsx' or not zipfile.is_zipfile(path):
        return {'ok': False, 'text': t('bad', code, reason=reason('not_xlsx', code))}
    try:
        import openpyxl
        book = openpyxl.load_workbook(str(path), data_only=False, keep_links=False)
    except Exception:
        return {'ok': False, 'text': t('bad', code, reason=reason('unreadable', code))}
    sheet = str(args['sheet'])
    match = next((ws for ws in book.worksheets if ws.title.casefold() == sheet.casefold()), None)
    if match is None:
        return {'ok': False, 'text': t('bad', code, reason=reason('no_sheet', code, sheet=sheet, sheets=', '.join(ws.title for ws in book.worksheets[:10]))), 'clarify': True}
    if match.sheet_state != 'visible':
        return {'ok': False, 'text': t('bad', code, reason=reason('hidden_sheet', code, sheet=match.title))}
    rng = str(args['range']).upper()
    try:
        min_col, min_row, max_col, max_row = _coords(rng)
    except Exception:
        min_col = min_row = max_col = max_row = 0
    if not (min_col and min_row) or not 2 <= max_col - min_col + 1 <= 1 + MAX_SERIES or not 2 <= max_row - min_row + 1 <= MAX_ROWS + 1:
        return {'ok': False, 'text': t('bad', code, reason=reason('bad_range', code, range=rng))}
    merged = [str(m) for m in match.merged_cells.ranges if not (m.max_col < min_col or m.min_col > max_col or m.max_row < min_row or m.min_row > max_row)]
    if merged:
        return {'ok': False, 'text': t('bad', code, reason=reason('merged', code, cells=', '.join(merged[:3])))}
    from openpyxl.utils import get_column_letter
    hidden = any(match.row_dimensions[r].hidden for r in range(min_row, max_row + 1)) or \
        any(match.column_dimensions[get_column_letter(c)].hidden for c in range(min_col, max_col + 1))
    if hidden:
        return {'ok': False, 'text': t('bad', code, reason=reason('hidden', code))}
    rows = [list(r) for r in match.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)]
    headers = []
    for cell in rows[0]:
        if cell.data_type == 'f':
            return {'ok': False, 'text': t('bad', code, reason=reason('formula', code, cell=cell.coordinate))}
        if cell.value is None or not str(cell.value).strip():
            return {'ok': False, 'text': t('bad', code, reason=reason('header', code, cell=cell.coordinate)), 'clarify': True}
        headers.append(str(cell.value).strip())
    data = []
    for row in rows[1:]:
        label = row[0]
        if label.data_type == 'f':
            return {'ok': False, 'text': t('bad', code, reason=reason('formula', code, cell=label.coordinate))}
        if label.value is None or not str(label.value).strip():
            return {'ok': False, 'text': t('bad', code, reason=reason('label', code, cell=label.coordinate))}
        values = []
        for cell in row[1:]:
            if cell.data_type == 'f':
                return {'ok': False, 'text': t('bad', code, reason=reason('formula', code, cell=cell.coordinate))}
            if cell.value is None:
                return {'ok': False, 'text': t('bad', code, reason=reason('empty', code, cell=cell.coordinate))}
            if isinstance(cell.value, bool) or not isinstance(cell.value, (int, float)) or not math.isfinite(float(cell.value)):
                return {'ok': False, 'text': t('bad', code, reason=reason('not_number', code, cell=cell.coordinate, value=str(cell.value)[:20]))}
            values.append(cell.value)
        data.append([label.value] + values)
    kind = 'bar' if str(args['chart']) == 'bar' else 'line'
    suffix = {'ru': 'диаграмма', 'en': 'chart', 'de': 'Diagramm', 'fr': 'graphique', 'es': 'gráfico', 'it': 'grafico', 'pt': 'gráfico'}.get(code, 'chart')
    destination = str(args.get('destination') or '')
    if not destination:
        from jarvis.runtime.office178 import _safe_name, _unique
        destination = str(_unique(path.parent, _safe_name(f'{path.stem} — {suffix}', 'chart'), '.xlsx'))
    kinds = _TEXT['kind'][_ORDER.index(code) if code in _ORDER else 1]
    sample = '; '.join(f'{r[0]}: ' + ', '.join(str(v) for v in r[1:]) for r in data[:3])
    text = t('preview', code, source=str(path), sheet=match.title, range=rng, category=headers[0], series=', '.join(f'«{h}»' for h in headers[1:]),
             rows=len(data), sample=sample, kind=kinds[0] if kind == 'bar' else kinds[1], destination=destination)
    return {'ok': True, 'text': text, 'path': str(path), 'sheet': match.title, 'range': rng, 'headers': headers, 'data': data,
            'chart': kind, 'destination': destination, 'fingerprint': fingerprint(path)}


def create(config: Any, args: dict[str, Any], code: str) -> tuple[bool, str, dict[str, Any]]:
    """Write the NEW workbook after re-checking the exact previewed source version."""
    spec = preview(config, dict(args), code)
    if not spec.get('ok'):
        return False, spec['text'], {'outcome': 'invalid'}
    if args.get('fingerprint') and spec['fingerprint'] != args['fingerprint']:
        return False, t('changed', code), {'outcome': 'source_changed'}
    destination = Path(spec['destination'])
    if destination.exists() or destination.suffix.casefold() != '.xlsx' or destination.parent != Path(spec['path']).parent:
        return False, t('changed', code), {'outcome': 'destination_changed'}
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, Reference
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = spec['sheet'][:31]
    sheet.append(spec['headers'])
    for row in spec['data']:
        sheet.append(row)
    rows = len(spec['data']) + 1
    chart = BarChart() if spec['chart'] == 'bar' else LineChart()
    if spec['chart'] == 'bar':
        chart.type = 'col'
    chart.title = ', '.join(spec['headers'][1:])
    chart.add_data(Reference(sheet, min_col=2, max_col=len(spec['headers']), min_row=1, max_row=rows), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=rows))
    sheet.add_chart(chart, f'{openpyxl.utils.get_column_letter(len(spec["headers"]) + 2)}2')
    temporary = destination.with_name('.~neo-' + uuid.uuid4().hex[:8] + '.xlsx')
    try:
        book.save(str(temporary))
        with zipfile.ZipFile(temporary) as archive:
            if not any(n.startswith('xl/charts/chart') for n in archive.namelist()):
                raise OSError('chart part missing')
        try:
            os.link(temporary, destination)            # fails if the name appeared meanwhile: never overwrite
            temporary.unlink()
        except (OSError, NotImplementedError, AttributeError):
            if destination.exists():
                raise FileExistsError(str(destination))
            os.replace(temporary, destination)
    except Exception:
        temporary.unlink(missing_ok=True)      # only our own temporary file; an existing file is never touched
        return False, t('bad', code, reason=reason('unreadable', code)), {'outcome': 'write_failed'}
    return True, t('done', code, destination=str(destination)), {'outcome': 'created', 'path': str(destination), 'source': spec['path']}
