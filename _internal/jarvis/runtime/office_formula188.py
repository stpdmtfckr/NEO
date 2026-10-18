"""NEO 188 NEO-OFFICE-002 / TZ-OFFICE-FORMULA-REFERENCE-COPY-01.

«В файле Продажи.xlsx скопируй формулу из C3 в D4» moves the formula like
Excel's copy/paste: relative parts of A1 references shift by the offset,
``$``-anchored parts stay (=A1+$A$1+$A1+A$1 from C3 to D4 → =B2+$A$1+$A2+B$1).
String literals, numbers, operators and function names are never changed.

Narrow on purpose: a saved .xlsx, one sheet, a source cell/rectangle of
formulas and a non-overlapping, EMPTY destination of the same shape. Sheet or
workbook references, structured/whole-column/R1C1 references, or a shift past
the sheet edge are refused (no #REF! is written). A workbook with charts,
pictures, pivot tables, macros or external links is not rewritten, because the
library would lose those parts. Preview first; the exact source version is
checked again before writing; a backup is kept for «отмени»; the written
formula strings are verified. Nothing is recalculated and Excel is not run.
"""
from __future__ import annotations

import os
import re
import uuid
import zipfile
from functools import lru_cache
from pathlib import Path
from typing import Any

from jarvis.runtime.natural_commands import LANGUAGES

MAX_CELLS = 400
MAX_COL, MAX_ROW = 16384, 1048576
_REF = re.compile(r'(?<![A-Za-z0-9_$.])(\$?)([A-Za-z]{1,3})(\$?)([0-9]{1,7})(?![A-Za-z0-9_(])')
_UNSUPPORTED = re.compile(r'!|\[|\bR\d*C\d*\b|(?<![A-Za-z0-9_$])\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}(?![0-9])|(?<![A-Za-z0-9_$])\$?\d+:\$?\d+(?![0-9])', re.I)
_RISKY_PARTS = ('xl/charts/', 'xl/drawings/', 'xl/media/', 'xl/pivotTables/', 'xl/pivotCache/', 'xl/vbaProject', 'xl/externalLinks/', 'xl/slicers/')


def _col_num(letters: str) -> int:
    value = 0
    for char in letters.upper():
        value = value * 26 + (ord(char) - 64)
    return value


def _col_name(number: int) -> str:
    out = ''
    while number:
        number, rest = divmod(number - 1, 26)
        out = chr(65 + rest) + out
    return out


def _segments(formula: str):
    """(is_string_literal, text) pieces: refs are shifted only outside "…"."""
    pieces, buffer, inside, i = [], '', False, 0
    while i < len(formula):
        char = formula[i]
        if char == '"':
            if inside and i + 1 < len(formula) and formula[i + 1] == '"':
                buffer += '""'
                i += 2
                continue
            if not inside:
                pieces.append((False, buffer))
                buffer, inside = '"', True
            else:
                pieces.append((True, buffer + '"'))
                buffer, inside = '', False
            i += 1
            continue
        buffer += char
        i += 1
    if inside:
        raise ValueError('unterminated string literal')
    pieces.append((False, buffer))
    return pieces


def shift(formula: str, rows: int, cols: int) -> str:
    """The formula as it reads after a copy by (rows, cols); ValueError when it cannot be shifted safely."""
    if not str(formula).startswith('='):
        raise ValueError('not a formula')
    out = []
    for literal, text in _segments(str(formula)):
        if literal:
            out.append(text)
            continue
        if _UNSUPPORTED.search(text):
            raise ValueError('unsupported reference')

        def move(match: re.Match) -> str:
            col_abs, col, row_abs, row = match.groups()
            col_n, row_n = _col_num(col), int(row)
            if not col_abs:
                col_n += cols
            if not row_abs:
                row_n += rows
            if not (1 <= col_n <= MAX_COL and 1 <= row_n <= MAX_ROW):
                raise ValueError('reference leaves the sheet')
            return f'{col_abs}{_col_name(col_n)}{row_abs}{row_n}'
        out.append(_REF.sub(move, text))
    return ''.join(out)


_CELL = r'\$?[A-Za-z]{1,3}\$?\d{1,7}'
_RNG = r'(?P<src>' + _CELL + r'(?:\s*:\s*' + _CELL + r')?)'
_DST = r'(?P<dst>' + _CELL + r'(?:\s*:\s*' + _CELL + r')?)'
_COPY = {
    'ru': r'(?:скопируй|скопировать|перенеси|копируй)\s+формул[уыа]?\s+(?:из\s+)?' + _RNG + r'\s+(?:в|на)\s+' + _DST,
    'en': r'copy\s+(?:the\s+)?formulas?\s+(?:from\s+)?' + _RNG + r'\s+(?:to|into)\s+' + _DST,
    'de': r'kopiere\s+(?:die\s+)?formeln?\s+(?:von|aus)\s+' + _RNG + r'\s+(?:nach|in)\s+' + _DST,
    'fr': r'copie\s+(?:la\s+|les\s+)?formules?\s+(?:de\s+)?' + _RNG + r'\s+(?:vers|dans|à)\s+' + _DST,
    'es': r'copia\s+(?:la\s+|las\s+)?fórmulas?\s+(?:de\s+)?' + _RNG + r'\s+(?:a|en)\s+' + _DST,
    'it': r'copia\s+(?:la\s+|le\s+)?formul[ae]\s+(?:da\s+)?' + _RNG + r'\s+(?:a|in)\s+' + _DST,
    'pt': r'copia\s+(?:a\s+|as\s+)?fórmulas?\s+(?:de\s+)?' + _RNG + r'\s+(?:para|em)\s+' + _DST,
}
_FILE = re.compile(r'(?:«(?P<q>[^»]+?\.xlsx)»|"(?P<d>[^"]+?\.xlsx)"|(?P<p>[^\s,;:«»"]+?\.xlsx))', re.I)


@lru_cache(maxsize=7)
def _pattern(code: str):
    return re.compile(_COPY[code], re.I)


def parse(text: str, code: str) -> tuple[str, dict] | None:
    if code not in LANGUAGES:
        return None
    value = str(text or '')
    match = _pattern(code).search(value)
    if not match:
        return None
    args: dict[str, Any] = {'action': 'copy_formula', 'source_range': re.sub(r'\s+', '', match.group('src')).upper(),
                            'target': re.sub(r'\s+', '', match.group('dst')).upper()}
    file = _FILE.search(value)
    if file:
        args['name'] = (file.group('q') or file.group('d') or file.group('p')).strip()
    from jarvis.runtime.office_chart188 import _SHEET, _unquote
    sheet = re.search(_SHEET[code], value, re.I)
    if sheet:
        args['sheet'] = _unquote(sheet.group('s'))
    return 'office', args


_ORDER = ('ru', 'en', 'de', 'fr', 'es', 'it', 'pt')
_TEXT = {
    'need': ('Для копирования формул нужен файл .xlsx: «в файле Отчёт.xlsx скопируй формулу из C3 в D4».',
             'Copying formulas needs the .xlsx file: “copy the formula from C3 to D4 in Report.xlsx”.',
             'Zum Kopieren von Formeln brauche ich die .xlsx-Datei: „kopiere die Formel von C3 nach D4 in Bericht.xlsx“.',
             'Pour copier des formules, il faut le fichier .xlsx : « copie la formule de C3 vers D4 dans Rapport.xlsx ».',
             'Para copiar fórmulas hace falta el archivo .xlsx: «copia la fórmula de C3 a D4 en Informe.xlsx».',
             'Per copiare formule serve il file .xlsx: «copia la formula da C3 a D4 in Report.xlsx».',
             'Para copiar fórmulas é preciso o ficheiro .xlsx: «copia a fórmula de C3 para D4 em Relatorio.xlsx».'),
    'bad': ('Формулы не копирую: {reason}. Ничего не записано.', 'I am not copying the formulas: {reason}. Nothing was written.',
            'Ich kopiere die Formeln nicht: {reason}. Nichts wurde geschrieben.', 'Je ne copie pas les formules : {reason}. Rien n’a été écrit.',
            'No copio las fórmulas: {reason}. No se escribió nada.', 'Non copio le formule: {reason}. Non è stato scritto nulla.',
            'Não copio as fórmulas: {reason}. Nada foi escrito.'),
    'preview': ('Копирование формул в {source}, лист «{sheet}»:\n{plan}\nДругие ячейки не меняются; копия книги сохранится для отмены. Выполнить?',
                'Copying formulas in {source}, sheet “{sheet}”:\n{plan}\nNo other cell changes; a backup is kept for undo. Go ahead?',
                'Formeln kopieren in {source}, Blatt „{sheet}“:\n{plan}\nAndere Zellen bleiben gleich; eine Sicherung wird behalten. Ausführen?',
                'Copie de formules dans {source}, feuille « {sheet} » :\n{plan}\nAucune autre cellule ne change ; une sauvegarde est conservée. Exécuter ?',
                'Copia de fórmulas en {source}, hoja «{sheet}»:\n{plan}\nNo cambia ninguna otra celda; se guarda una copia para deshacer. ¿Lo hago?',
                'Copia di formule in {source}, foglio «{sheet}»:\n{plan}\nNessun’altra cella cambia; si conserva una copia per annullare. Procedo?',
                'Cópia de fórmulas em {source}, folha «{sheet}»:\n{plan}\nNenhuma outra célula muda; fica uma cópia para desfazer. Faço?'),
    'done': ('Скопировал формулы ({count}) в {source}; записанные формулы совпадают с планом. Пересчёт значений выполнит Excel при открытии.',
             'Copied {count} formula(s) in {source}; the written formulas match the plan. Excel recalculates the values when it opens the file.',
             '{count} Formel(n) in {source} kopiert; die geschriebenen Formeln entsprechen dem Plan. Excel berechnet die Werte beim Öffnen neu.',
             '{count} formule(s) copiée(s) dans {source} ; les formules écrites correspondent au plan. Excel recalculera à l’ouverture.',
             'Copié {count} fórmula(s) en {source}; las fórmulas escritas coinciden con el plan. Excel recalculará al abrir el archivo.',
             'Copiate {count} formule in {source}; le formule scritte corrispondono al piano. Excel ricalcola all’apertura.',
             'Copiei {count} fórmula(s) em {source}; as fórmulas escritas correspondem ao plano. O Excel recalcula ao abrir.'),
    'changed': ('Файл изменился после предпросмотра — ничего не записал. Повторите команду для нового предпросмотра.',
                'The file changed after the preview — nothing written. Repeat the request for a new preview.',
                'Die Datei hat sich nach der Vorschau geändert – nichts geschrieben.', 'Le fichier a changé après l’aperçu : rien n’a été écrit.',
                'El archivo cambió tras la vista previa: no escribí nada.', 'Il file è cambiato dopo l’anteprima: non ho scritto nulla.',
                'O ficheiro mudou depois da pré-visualização: não escrevi nada.'),
}
_REASONS = {
    'shape': ('назначение должно быть одной ячейкой или диапазоном того же размера', 'the destination must be one cell or a range of the same size',
              'das Ziel muss eine Zelle oder ein gleich großer Bereich sein', 'la destination doit être une cellule ou une plage de même taille',
              'el destino debe ser una celda o un rango del mismo tamaño', 'la destinazione deve essere una cella o un intervallo della stessa dimensione',
              'o destino deve ser uma célula ou um intervalo do mesmo tamanho'),
    'overlap': ('источник и назначение пересекаются', 'source and destination overlap', 'Quelle und Ziel überschneiden sich', 'la source et la destination se chevauchent',
                'el origen y el destino se solapan', 'origine e destinazione si sovrappongono', 'a origem e o destino sobrepõem-se'),
    'occupied': ('ячейка {cell} в назначении не пуста', 'destination cell {cell} is not empty', 'Zielzelle {cell} ist nicht leer', 'la cellule de destination {cell} n’est pas vide',
                 'la celda de destino {cell} no está vacía', 'la cella di destinazione {cell} non è vuota', 'a célula de destino {cell} não está vazia'),
    'not_formula': ('в ячейке {cell} не формула', 'cell {cell} is not a formula', 'Zelle {cell} ist keine Formel', 'la cellule {cell} n’est pas une formule',
                    'la celda {cell} no es una fórmula', 'la cella {cell} non è una formula', 'a célula {cell} não é uma fórmula'),
    'unsupported': ('формулу в {cell} нельзя перенести безопасно (ссылка на другой лист или книгу, целый столбец, R1C1 или выход за край листа)',
                    'the formula in {cell} cannot be moved safely (a reference to another sheet or workbook, a whole column, R1C1 or a shift past the sheet edge)',
                    'die Formel in {cell} lässt sich nicht sicher verschieben (Bezug auf anderes Blatt/Mappe, ganze Spalte, Z1S1 oder über den Rand)',
                    'la formule en {cell} ne peut pas être déplacée sans risque (autre feuille ou classeur, colonne entière, L1C1 ou hors de la feuille)',
                    'la fórmula de {cell} no se puede mover con seguridad (otra hoja o libro, columna entera, F1C1 o fuera de la hoja)',
                    'la formula in {cell} non si può spostare in sicurezza (altro foglio o cartella, colonna intera, R1C1 o fuori dal foglio)',
                    'a fórmula em {cell} não pode ser movida com segurança (outra folha ou livro, coluna inteira, L1C1 ou fora da folha)'),
    'rich': ('книга содержит диаграммы, рисунки, сводные таблицы, макросы или внешние связи — при перезаписи они могли бы потеряться',
             'the workbook contains charts, pictures, pivot tables, macros or external links that could be lost when rewriting it',
             'die Mappe enthält Diagramme, Bilder, Pivot-Tabellen, Makros oder externe Verknüpfungen, die beim Neuschreiben verloren gehen könnten',
             'le classeur contient des graphiques, images, tableaux croisés, macros ou liens externes qui pourraient être perdus',
             'el libro tiene gráficos, imágenes, tablas dinámicas, macros o vínculos externos que podrían perderse',
             'la cartella contiene grafici, immagini, tabelle pivot, macro o collegamenti esterni che potrebbero andare persi',
             'o livro tem gráficos, imagens, tabelas dinâmicas, macros ou ligações externas que se poderiam perder'),
    'too_many': ('слишком большой диапазон (больше 400 ячеек)', 'the range is too large (more than 400 cells)', 'der Bereich ist zu groß (über 400 Zellen)',
                 'la plage est trop grande (plus de 400 cellules)', 'el rango es demasiado grande (más de 400 celdas)', 'l’intervallo è troppo grande (oltre 400 celle)',
                 'o intervalo é demasiado grande (mais de 400 células)'),
    'sheet': ('укажите лист: в книге их несколько ({sheets})', 'name the sheet: the workbook has several ({sheets})', 'nenne das Blatt: Die Mappe hat mehrere ({sheets})',
              'indiquez la feuille : le classeur en a plusieurs ({sheets})', 'indica la hoja: el libro tiene varias ({sheets})', 'indica il foglio: la cartella ne ha più di uno ({sheets})',
              'indica a folha: o livro tem várias ({sheets})'),
    'verify': ('после записи формулы не совпали с планом — изменение отменено', 'the written formulas did not match the plan — the change was undone',
               'die geschriebenen Formeln stimmen nicht mit dem Plan überein – zurückgenommen', 'les formules écrites ne correspondent pas au plan : modification annulée',
               'las fórmulas escritas no coinciden con el plan: cambio deshecho', 'le formule scritte non corrispondono al piano: modifica annullata',
               'as fórmulas escritas não coincidem com o plano: alteração desfeita'),
}


def t(key: str, code: str, **values) -> str:
    return _TEXT[key][_ORDER.index(code) if code in _ORDER else 1].format(**values)


def reason(key: str, code: str, **values) -> str:
    return _REASONS[key][_ORDER.index(code) if code in _ORDER else 1].format(**values)


def _bounds(rng: str) -> tuple[int, int, int, int]:
    from openpyxl.utils.cell import range_boundaries
    return range_boundaries(rng.replace('$', ''))


def preview(config: Any, args: dict[str, Any], code: str) -> dict[str, Any]:
    from jarvis.runtime import office_chart188 as chart
    if not str(args.get('name', '') or '').strip():
        return {'ok': False, 'text': t('need', code), 'clarify': True}
    candidates = [args['path']] if args.get('path') else chart._candidates(str(args['name']))
    if not candidates:
        return {'ok': False, 'text': t('bad', code, reason=chart.reason('missing_file', code, name=args['name']))}
    if len(candidates) > 1:
        return {'ok': False, 'text': t('bad', code, reason=chart.reason('ambiguous', code, name=args['name'], paths='; '.join(candidates[:5]))), 'clarify': True}
    path = Path(candidates[0])
    if path.suffix.casefold() != '.xlsx' or not zipfile.is_zipfile(path):
        return {'ok': False, 'text': t('bad', code, reason=chart.reason('not_xlsx', code))}
    with zipfile.ZipFile(path) as archive:
        if any(name.startswith(_RISKY_PARTS) for name in archive.namelist()):
            return {'ok': False, 'text': t('bad', code, reason=reason('rich', code))}
    try:
        import openpyxl
        book = openpyxl.load_workbook(str(path), data_only=False, keep_links=False)
    except Exception:
        return {'ok': False, 'text': t('bad', code, reason=chart.reason('unreadable', code))}
    sheet_name = str(args.get('sheet', '') or '')
    if sheet_name:
        sheet = next((ws for ws in book.worksheets if ws.title.casefold() == sheet_name.casefold()), None)
        if sheet is None:
            return {'ok': False, 'text': t('bad', code, reason=chart.reason('no_sheet', code, sheet=sheet_name, sheets=', '.join(ws.title for ws in book.worksheets[:10]))), 'clarify': True}
    elif len(book.worksheets) == 1:
        sheet = book.worksheets[0]
    else:
        return {'ok': False, 'text': t('bad', code, reason=reason('sheet', code, sheets=', '.join(ws.title for ws in book.worksheets[:10]))), 'clarify': True}
    try:
        s_col1, s_row1, s_col2, s_row2 = _bounds(str(args['source_range']))
        d_col1, d_row1, d_col2, d_row2 = _bounds(str(args['target']))
    except Exception:
        return {'ok': False, 'text': t('bad', code, reason=reason('shape', code))}
    height, width = s_row2 - s_row1 + 1, s_col2 - s_col1 + 1
    if height * width > MAX_CELLS:
        return {'ok': False, 'text': t('bad', code, reason=reason('too_many', code))}
    if (d_col1, d_row1) != (d_col2, d_row2) and (d_row2 - d_row1 + 1, d_col2 - d_col1 + 1) != (height, width):
        return {'ok': False, 'text': t('bad', code, reason=reason('shape', code))}
    d_col2, d_row2 = d_col1 + width - 1, d_row1 + height - 1
    if not (d_col2 < s_col1 or d_col1 > s_col2 or d_row2 < s_row1 or d_row1 > s_row2):
        return {'ok': False, 'text': t('bad', code, reason=reason('overlap', code))}
    drow, dcol = d_row1 - s_row1, d_col1 - s_col1
    plan = []
    for r in range(height):
        for c in range(width):
            source = sheet.cell(row=s_row1 + r, column=s_col1 + c)
            target = sheet.cell(row=d_row1 + r, column=d_col1 + c)
            if target.value is not None:
                return {'ok': False, 'text': t('bad', code, reason=reason('occupied', code, cell=target.coordinate))}
            if source.value is None:
                continue
            if source.data_type != 'f' or not str(source.value).startswith('='):
                return {'ok': False, 'text': t('bad', code, reason=reason('not_formula', code, cell=source.coordinate))}
            try:
                moved = shift(str(source.value), drow, dcol)
            except ValueError:
                return {'ok': False, 'text': t('bad', code, reason=reason('unsupported', code, cell=source.coordinate))}
            plan.append({'from': source.coordinate, 'to': target.coordinate, 'formula': str(source.value), 'moved': moved})
    if not plan:
        return {'ok': False, 'text': t('bad', code, reason=reason('not_formula', code, cell=str(args['source_range'])))}
    lines = '\n'.join(f"• {p['from']}: {p['formula']} → {p['to']}: {p['moved']}" for p in plan[:20])
    if len(plan) > 20:
        lines += f'\n• … +{len(plan) - 20}'
    from jarvis.runtime.office_chart188 import fingerprint
    return {'ok': True, 'text': t('preview', code, source=str(path), sheet=sheet.title, plan=lines), 'path': str(path), 'sheet': sheet.title,
            'plan': plan, 'fingerprint': fingerprint(path)}


def apply(config: Any, args: dict[str, Any], code: str) -> tuple[bool, str, dict[str, Any]]:
    spec = preview(config, dict(args), code)
    if not spec.get('ok'):
        return False, spec['text'], {'outcome': 'invalid'}
    if args.get('fingerprint') and spec['fingerprint'] != args['fingerprint']:
        return False, t('changed', code), {'outcome': 'source_changed'}
    import openpyxl
    from jarvis.runtime import office178
    path = Path(spec['path'])
    book = openpyxl.load_workbook(str(path), data_only=False, keep_links=False)
    sheet = book[spec['sheet']]
    for step in spec['plan']:
        sheet[step['to']] = step['moved']
    temporary = path.with_name('.~neo-' + uuid.uuid4().hex[:8] + '.xlsx')
    try:
        book.save(str(temporary))
        check = openpyxl.load_workbook(str(temporary), data_only=False)[spec['sheet']]
        if any(str(check[step['to']].value) != step['moved'] for step in spec['plan']) or \
                any(str(check[step['from']].value) != step['formula'] for step in spec['plan']):
            temporary.unlink(missing_ok=True)
            return False, t('bad', code, reason=reason('verify', code)), {'outcome': 'verify_failed'}
        from jarvis.runtime.office_chart188 import fingerprint
        if fingerprint(path) != spec['fingerprint']:
            temporary.unlink(missing_ok=True)       # changed while we were writing: nothing is replaced
            return False, t('changed', code), {'outcome': 'source_changed'}
        office178._backup(config, path)            # the existing «отмени» copy of the original
        os.replace(temporary, path)
    except Exception:
        temporary.unlink(missing_ok=True)
        return False, t('bad', code, reason=office178_reason(code)), {'outcome': 'write_failed'}
    office178.remember(path, 'xlsx')
    return True, t('done', code, count=len(spec['plan']), source=str(path)), {'outcome': 'copied', 'path': str(path), 'cells': [p['to'] for p in spec['plan']]}


def office178_reason(code: str) -> str:
    from jarvis.runtime.office_chart188 import reason as chart_reason
    return chart_reason('unreadable', code)
