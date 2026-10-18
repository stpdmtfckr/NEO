"""NEO 188 NEO-OFFICE-002 / TZ-OFFICE-FORMULA-REFERENCE-COPY-01 (synthetic workbooks)."""
import zipfile
from types import SimpleNamespace

import pytest

openpyxl = pytest.importorskip('openpyxl')

from jarvis.runtime import office178
from jarvis.runtime import office_formula188 as formula


def test_ofc01_relative_absolute_and_mixed_references():
    assert formula.shift('=A1+$A$1+$A1+A$1', 1, 1) == '=B2+$A$1+$A2+B$1'


def test_ofc02_string_literals_numbers_and_functions_are_untouched():
    assert formula.shift('=IF(A1>10,"A1 ""B2""",SUM(A1:A3)*2)', 1, 0) == '=IF(A2>10,"A1 ""B2""",SUM(A2:A4)*2)'
    assert formula.shift('=LOG10(A1)', 0, 1) == '=LOG10(B1)'


@pytest.mark.parametrize('bad', ['=Sheet2!A1', "='My sheet'!A1", '=[Book.xlsx]S!A1', '=SUM(A:A)', '=SUM(1:1)', '=R1C1', '=Table1[Col]'])
def test_ofc04_unsupported_references_are_refused(bad):
    with pytest.raises(ValueError):
        formula.shift(bad, 1, 1)


def test_ofc04_shift_past_the_edge_is_refused():
    with pytest.raises(ValueError):
        formula.shift('=A1', -1, 0)


@pytest.fixture
def book(tmp_path, monkeypatch):
    monkeypatch.setattr(office178, 'user_folders', lambda: [tmp_path])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Лист1'
    ws['A1'], ws['A2'], ws['B1'] = 1, 2, 3
    ws['C3'] = '=A1+$A$1+$A1+A$1'
    ws['C4'] = '=A2*2'
    ws['F9'] = 'не трогать'
    path = tmp_path / 'Отчёт.xlsx'
    wb.save(path)
    config = SimpleNamespace(paths=SimpleNamespace(root=tmp_path, cache=tmp_path / 'cache'), get=lambda k, d=None: d)
    return path, config


def args(**extra):
    return {'action': 'copy_formula', 'name': 'Отчёт.xlsx', 'source_range': 'C3:C4', 'target': 'D4', **extra}


def test_ofc03_rectangle_copied_cell_by_cell_and_nothing_else_changes(book):
    path, config = book
    spec = formula.preview(config, args(), 'ru')
    assert spec['ok'] and 'C3: =A1+$A$1+$A1+A$1 → D4: =B2+$A$1+$A2+B$1' in spec['text'] and 'D5: =B3*2' in spec['text']
    ok, text, data = formula.apply(config, args(fingerprint=spec['fingerprint'], path=spec['path'], sheet=spec['sheet']), 'ru')
    assert ok and data['cells'] == ['D4', 'D5'] and 'совпадают с планом' in text
    ws = openpyxl.load_workbook(path)['Лист1']
    assert ws['D4'].value == '=B2+$A$1+$A2+B$1' and ws['D5'].value == '=B3*2'
    assert ws['C3'].value == '=A1+$A$1+$A1+A$1' and ws['F9'].value == 'не трогать' and ws['D6'].value is None
    assert any((config.paths.cache / 'office_undo').iterdir())        # the undo copy exists


def test_ofc04_occupied_or_overlapping_destination_writes_nothing(book):
    path, config = book
    before = path.read_bytes()
    assert not formula.preview(config, args(target='F9'), 'ru')['ok']
    assert not formula.preview(config, args(target='C4'), 'ru')['ok']
    assert not formula.preview(config, args(target='D4:E5'), 'ru')['ok']
    assert path.read_bytes() == before


def test_ofc05_changed_file_after_preview_is_not_written(book):
    path, config = book
    spec = formula.preview(config, args(), 'ru')
    wb = openpyxl.load_workbook(path)
    wb.active['A1'] = 5
    wb.save(path)
    ok, _, data = formula.apply(config, args(fingerprint=spec['fingerprint'], path=spec['path'], sheet=spec['sheet']), 'ru')
    assert not ok and data['outcome'] == 'source_changed'
    assert openpyxl.load_workbook(path).active['D4'].value is None


def test_rich_workbook_is_not_rewritten(book, tmp_path):
    path, config = book
    from openpyxl.chart import BarChart, Reference
    wb = openpyxl.load_workbook(path)
    chart = BarChart()
    chart.add_data(Reference(wb.active, min_col=1, min_row=1, max_row=2))
    wb.active.add_chart(chart, 'H2')
    wb.save(path)
    spec = formula.preview(config, args(), 'ru')
    assert not spec['ok'] and 'диаграммы' in spec['text']


def test_values_are_not_formulas(book):
    path, config = book
    assert not formula.preview(config, args(source_range='A1:A2', target='H1'), 'ru')['ok']


@pytest.mark.parametrize('code,phrase', [
    ('ru', 'в файле Отчёт.xlsx скопируй формулу из C3 в D4'), ('en', 'copy the formula from C3 to D4 in Report.xlsx'),
    ('de', 'kopiere die Formel von C3 nach D4 in Bericht.xlsx'), ('fr', 'copie la formule de C3 vers D4 dans Rapport.xlsx'),
    ('es', 'copia la fórmula de C3 a D4 en Informe.xlsx'), ('it', 'copia la formula da C3 a D4 in Report.xlsx'),
    ('pt', 'copia a fórmula de C3 para D4 em Relatorio.xlsx'),
])
def test_phrases_in_seven_languages(code, phrase):
    tool, parsed = formula.parse(phrase, code)
    assert tool == 'office' and parsed == {'action': 'copy_formula', 'source_range': 'C3', 'target': 'D4', 'name': parsed['name']}
    assert parsed['name'].endswith('.xlsx')
